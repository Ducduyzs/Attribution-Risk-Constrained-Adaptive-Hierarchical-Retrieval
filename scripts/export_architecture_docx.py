from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "EDAHR_Kien_truc_pipeline_va_nen_tang_kien_thuc.docx"

NAVY = "17324D"
TEAL = "087E8B"
ORANGE = "E07A3F"
PALE_TEAL = "E8F3F4"
PALE_ORANGE = "FAEEE7"
LIGHT = "F3F6F8"
MID = "D7E0E6"
WHITE = "FFFFFF"
TEXT = RGBColor(35, 48, 58)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 30, NAVY),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, TEAL),
        ("Heading 3", 11, ORANGE),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(12)
    styles["Heading 1"].paragraph_format.page_break_before = True

    footer = section.footer.paragraphs[0]
    footer.add_run("EDAHR | Tài liệu kiến trúc kỹ thuật")
    footer.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    footer.runs[0].font.size = Pt(9)
    footer.add_run("\t")
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.0))
    add_page_number(footer)

    properties = document.core_properties
    properties.title = "Kiến trúc pipeline và nền tảng kiến thức EDAHR"
    properties.subject = "Attribution-risk-aware adaptive hierarchical retrieval"
    properties.author = "Nhóm phát triển EDAHR"
    properties.keywords = "RAG, hierarchical retrieval, attribution risk, scientific QA"


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(38)
    run = paragraph.add_run("TÀI LIỆU KIẾN TRÚC KỸ THUẬT")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(ORANGE)

    title = document.add_paragraph(style="Title")
    title.add_run("EDAHR\n")
    subtitle = title.add_run("Pipeline và nền tảng kiến thức xây dựng hệ thống")
    subtitle.font.size = Pt(22)
    subtitle.font.color.rgb = RGBColor.from_string(TEAL)

    line = document.add_table(rows=1, cols=2)
    line.alignment = WD_TABLE_ALIGNMENT.LEFT
    line.columns[0].width = Cm(3.0)
    line.columns[1].width = Cm(13.0)
    set_cell_shading(line.cell(0, 0), ORANGE)
    set_cell_shading(line.cell(0, 1), NAVY)
    for cell in line.rows[0].cells:
        cell.height = Cm(0.2)
        cell.text = ""

    document.add_paragraph()
    full_name = document.add_paragraph()
    full_name.add_run(
        "Attribution-Risk-Aware Adaptive Hierarchical Retrieval\n"
        "for Verifiable Scientific Document Question Answering"
    ).bold = True
    full_name.runs[0].font.size = Pt(16)
    full_name.runs[0].font.color.rgb = RGBColor.from_string(NAVY)

    add_callout(
        document,
        "Mục tiêu",
        "Trả lời câu hỏi trên tài liệu khoa học với bằng chứng truy nguyên đến đoạn leaf cụ thể; "
        "mở rộng ngữ cảnh theo phân cấp chỉ khi lợi ích lớn hơn rủi ro sai quy kết và chi phí token.",
        PALE_TEAL,
    )

    document.add_paragraph("Phiên bản tài liệu: 1.0")
    document.add_paragraph(f"Ngày chốt nội dung: {date(2026, 9, 1).strftime('%d/%m/%Y')}")
    document.add_paragraph("Nguồn: mã nguồn và báo cáo V7 trong repository")

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(55)
    run = note.add_run(
        "Phạm vi: kiến trúc hiện hành, thuật toán cốt lõi, pipeline huấn luyện, "
        "đánh giá, vận hành và các giới hạn nghiên cứu đang mở."
    )
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(TEAL)
    document.add_page_break()


def add_toc(document: Document) -> None:
    document.add_heading("Mục lục", level=1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Cập nhật trường mục lục trong Microsoft Word (Ctrl+A, F9)."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)
    document.add_page_break()


def add_callout(document: Document, title: str, text: str, fill: str = LIGHT) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, TEAL, "8")
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title + "\n")
    title_run.bold = True
    title_run.font.name = "Aptos"
    title_run.font.color.rgb = RGBColor.from_string(TEAL)
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[tuple[str, str]]) -> None:
    for title, text in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(title + ": ")
        run.bold = True
        paragraph.add_run(text)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, WHITE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.name = "Aptos"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell.paragraphs[0].add_run(str(value))
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph()


def add_pipeline_diagram(document: Document) -> None:
    stages = [
        ("01", "PDF khoa học", "Docling + layout/page"),
        ("02", "Phân cấp", "Document > Section > Parent > Child"),
        ("03", "Truy hồi leaf", "Dense + sparse + ColBERT"),
        ("04", "Rerank", "BGE cross-encoder"),
        ("05", "Mở rộng thích nghi", "Parent gate + section gate + rollback"),
        ("06", "Lắp ngữ cảnh", "Dedup + knapsack + diversity"),
        ("07", "Sinh claim", "Strict JSON + context IDs"),
        ("08", "Kiểm chứng leaf", "NLI + safety guards"),
        ("09", "Kết quả", "Claim + evidence + trace + metrics"),
    ]
    table = document.add_table(rows=len(stages), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (number, title, detail) in enumerate(stages):
        cells = table.rows[index].cells
        set_cell_shading(cells[0], ORANGE if index in {4, 7} else TEAL)
        set_cell_shading(cells[1], PALE_ORANGE if index in {4, 7} else PALE_TEAL)
        set_cell_shading(cells[2], WHITE)
        for cell in cells:
            set_cell_border(cell, WHITE, "12")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        number_run = cells[0].paragraphs[0].add_run(number)
        number_run.bold = True
        number_run.font.color.rgb = RGBColor.from_string(WHITE)
        number_run.font.size = Pt(12)
        title_run = cells[1].paragraphs[0].add_run(title)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor.from_string(NAVY)
        cells[2].paragraphs[0].add_run(detail)
    document.add_paragraph()


def add_formula(document: Document, formula: str, explanation: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_border(cell, MID)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(formula)
    run.bold = True
    run.font.name = "Cambria Math"
    run.font.size = Pt(12)
    detail = cell.add_paragraph(explanation)
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.runs[0].italic = True
    document.add_paragraph()


def add_chapter_overview(document: Document) -> None:
    document.add_heading("1. Tổng quan hệ thống", level=1)
    document.add_heading("1.1 Bài toán", level=2)
    document.add_paragraph(
        "EDAHR là hệ thống hỏi đáp trên PDF khoa học. Bài toán không dừng ở việc tìm ngữ cảnh "
        "có liên quan; mỗi mệnh đề trong câu trả lời còn phải truy nguyên được đến một đoạn leaf "
        "cụ thể, đúng tài liệu, đúng vùng trang và vượt qua kiểm chứng hỗ trợ."
    )
    add_callout(
        document,
        "Luận điểm trung tâm",
        "Ngữ cảnh rộng giúp tăng độ đầy đủ nhưng làm tăng nhiễu và rủi ro gán sai bằng chứng. "
        "Hệ thống vì thế học khi nào nên KEEP leaf, EXPAND lên parent, hoặc tiếp tục lên section/document.",
        PALE_ORANGE,
    )
    document.add_heading("1.2 Nguyên tắc thiết kế", level=2)
    add_bullets(document, [
        "Leaf-first retrieval: truy hồi ban đầu ở mức child để giữ độ đặc hiệu.",
        "Adaptive hierarchy: không mở rộng cố định; quyết định theo từng query và nhóm bằng chứng.",
        "Attribution before trust: claim chỉ tồn tại trong đáp án cuối khi có leaf hỗ trợ.",
        "Budget awareness: mọi lợi ích mở rộng được cân với chi phí token và độ nhiễu.",
        "Fail closed: citation sai schema, NLI label không rõ hoặc contradiction mạnh đều bị từ chối.",
        "Traceability: lưu hit, quyết định merge, rollback, context, generation thô, verification trace và metrics.",
    ])
    document.add_heading("1.3 Hai pipeline liên kết", level=2)
    add_table(document, ["Pipeline", "Mục đích", "Đầu ra"], [
        ["Online inference", "Trả lời một câu hỏi từ một hay nhiều PDF", "Verified claims + leaf evidence + diagnostics"],
        ["Offline learning", "Tạo nhãn phản thực và học hai gate mở rộng", "Parent/section checkpoint + metadata"],
        ["Benchmark", "Đo retrieval, answer, citation, risk, chi phí và thống kê", "JSON artifacts + reports + provenance"],
    ], [3.4, 7.0, 6.0])


def add_online_pipeline(document: Document) -> None:
    document.add_heading("2. Kiến trúc pipeline trực tuyến", level=1)
    document.add_paragraph(
        "Điểm điều phối chính là AdaptiveHierarchicalPipeline.answer(). Luồng dưới đây chạy tuần tự, "
        "có telemetry thời gian cho từng pha và trả về một Result có cấu trúc."
    )
    add_pipeline_diagram(document)

    document.add_heading("2.1 Phân loại câu hỏi", level=2)
    add_table(document, ["Loại query", "Tín hiệu", "Trần mở rộng"], [
        ["Factoid", "Hỏi sự kiện/giá trị cụ thể", "Parent"],
        ["Explanatory", "Why/how/mechanism/rationale", "Section"],
        ["Comparative", "Compare/versus/difference", "Document"],
        ["Global", "Overall/summary/main contribution", "Document"],
    ])
    document.add_paragraph(
        "Bộ phân loại hiện là luật từ khóa tiếng Anh và tiếng Việt. Đây là guardrail điều khiển không gian "
        "tìm kiếm, chưa phải một classifier học máy."
    )

    document.add_heading("2.2 Truy hồi đa biểu diễn", level=2)
    add_numbered(document, [
        ("Dense", "BGE-M3 tạo vector dense, chuẩn hóa L2; FAISS IndexFlatIP tìm láng giềng theo inner product."),
        ("Learned sparse", "Trọng số lexical của BGE-M3 được lưu trong postings; score là tích trọng số query-document theo token."),
        ("ColBERT", "Late interaction lấy max similarity theo từng query token rồi lấy trung bình."),
        ("Fusion", "Weighted fusion sau min-max normalization hoặc reciprocal-rank fusion (RRF)."),
    ])
    add_formula(
        document,
        "Score = w_dense·d + w_sparse·s + w_colbert·c",
        "Mặc định: 0,40 / 0,25 / 0,35; trọng số chỉ được chuẩn hóa trên các biểu diễn đang bật.",
    )

    document.add_heading("2.3 Rerank và cache điểm", level=2)
    document.add_paragraph(
        "Top candidate_k leaf được lấy từ index; top rerank_k được BGE reranker chấm lại theo cặp "
        "(query, passage). Reranker score trở thành tín hiệu relevance trực tiếp cho merge feature và "
        "được cache để tránh chấm lặp ancestor."
    )

    document.add_heading("2.4 Gate child → parent", level=2)
    document.add_paragraph(
        "Các child hit chung parent tạo thành candidate group. Hệ thống so sánh giữ tập member với thay "
        "bằng parent. Gate có thể là calibrated prior, checkpoint TorchScript hoặc mô hình scikit-learn."
    )
    add_formula(
        document,
        "Δ = U(parent) + w_gain·gain·U(parent) − U(members) − w_cost·Δtokens",
        "Prior nhận merge khi Δ ≥ margin. Learned policy nhận merge khi P(merge) ≥ threshold.",
    )
    add_callout(
        document,
        "Rollback guard",
        "Dù policy muốn merge, candidate bị rollback nếu reranker score của ancestor < rollback_ratio × "
        "best member score. Guard này ngăn context rộng làm mất độ đặc hiệu của bằng chứng tốt nhất.",
    )

    document.add_heading("2.5 Gate parent → section → document", level=2)
    document.add_paragraph(
        "expand_selection() lặp từng mức, dùng section policy độc lập. Vòng lặp dừng khi đạt trần theo "
        "query type, hết expansion_max_depth, không đủ member, không merge nào được chấp nhận, tổng gain "
        "dưới epsilon hoặc tổng token vượt headroom."
    )

    document.add_heading("2.6 Lắp context", level=2)
    add_numbered(document, [
        ("Khử trùng lặp", "Greedy Jaccard/containment loại block gần trùng."),
        ("Giới hạn pool", "Giữ tối đa final_context_k candidate có utility cao."),
        ("Knapsack 0/1", "Tối đa hóa tổng utility theo token bucket dưới ngân sách cứng."),
        ("Source diversity", "Giới hạn tỷ lệ block từ một nguồn; query comparative cố gắng có ít nhất hai nguồn."),
        ("Truncation", "Chỉ cắt theo biên câu như fallback cuối cùng."),
        ("Context IDs", "Gán C1, C2, ...; mỗi block vẫn mang evidence_child_ids để truy ngược leaf."),
    ])

    document.add_heading("2.7 Sinh và kiểm chứng claim", level=2)
    document.add_paragraph(
        "OpenAI dùng strict JSON Schema với enum context ID động. Gemini và Antigravity được parse và "
        "validate tương đương ở tầng ứng dụng. Claim answerable bắt buộc có citation dạng C1, không dùng [C1] "
        "hay chỉ số 1. Vi phạm được giữ trong validation_errors thay vì âm thầm xóa."
    )
    add_numbered(document, [
        ("Resolve citation", "Lấy các context block được claim trích dẫn."),
        ("Expand to leaves", "Chỉ xét child nằm trong evidence_ids của block được trích dẫn."),
        ("NLI", "Chấm entailment và contradiction cho từng cặp leaf → claim."),
        ("Safety", "Contradiction mạnh hoặc xung đột phủ định/số liệu đặt effective support về 0."),
        ("Lexical rescue", "Cứu paraphrase gần nguyên văn khi coverage đủ cao và không có safety conflict."),
        ("Sibling guard", "Leaf không nằm trong retrieval ban đầu phải vượt threshold cao hơn."),
        ("Select evidence", "Mặc định giữ top-1 leaf; trường hợp mơ hồ cũng chỉ giữ top-1 và ghi metric."),
    ])
    add_formula(
        document,
        "AR(S) = 1 − (1/|C|) Σ_c max_e P_entail(e ⇒ c)",
        "Attribution Risk cao khi claim không thể quy kết chắc chắn cho một leaf cụ thể.",
    )


def add_data_architecture(document: Document) -> None:
    document.add_heading("3. Kiến trúc dữ liệu và provenance", level=1)
    document.add_heading("3.1 Ingestion PDF", level=2)
    document.add_paragraph(
        "DoclingScientificLoader dùng Docling để đọc layout, heading, table và vùng trang. Kết quả trung gian "
        "là ScientificDocument gồm các DocumentSection; metadata giữ parser version, structure source, page "
        "spans và bounding boxes khi có."
    )
    document.add_heading("3.2 Cây phân cấp bốn mức", level=2)
    add_table(document, ["Mức", "Cách tạo", "Vai trò"], [
        ["Document", "Ghép toàn bộ section", "Ngữ cảnh tổng hợp toàn bài"],
        ["Section", "Theo cấu trúc Docling/dataset", "Đơn vị chủ đề lớn"],
        ["Parent", "Nhóm child có overlap cấu hình được", "Đơn vị mở rộng cục bộ"],
        ["Child", "Chunk khoảng 220 token, overlap một câu", "Đơn vị retrieval và attribution cuối"],
    ])
    document.add_heading("3.3 Chuỗi truy nguyên", level=2)
    add_table(document, ["Trường", "Ý nghĩa"], [
        ["node_id", "ID ổn định SHA-1 rút gọn từ document/level/position/text"],
        ["parent_id / child_ids", "Liên kết cây để đi lên và đi xuống"],
        ["evidence_child_ids", "Tập leaf tạo nên ancestor; cầu nối cho verification"],
        ["source", "Tên/đường dẫn tài liệu nguồn"],
        ["page_start / page_end", "Khoảng trang vật lý từ Docling"],
        ["char_start / char_end", "Offset ký tự để map chunk và đoạn gốc"],
        ["paragraph_ids", "ID paragraph QASPER giao với span child"],
        ["confidence", "Độ tin cậy cấu trúc/parser truyền xuống evidence"],
    ])
    add_callout(
        document,
        "Hai không gian đánh giá khác nhau",
        "Official QASPER Evidence F1 phải tính ở paragraph level. Leaf Attribution F1, harmful drift và rescue "
        "rate đo hành vi child-level của hệ thống. Không được gọi leaf citation F1 là official QASPER Evidence F1.",
        PALE_ORANGE,
    )
    document.add_heading("3.4 Mô hình kết quả", level=2)
    add_table(document, ["Đối tượng", "Nội dung"], [
        ["Hit", "Điểm dense, sparse, ColBERT, reranker và rank"],
        ["MergeDecision", "Features, probability, utility, accepted, rollback, token delta"],
        ["ContextBlock", "C-ID, node, level, text, provenance, leaf descendants, utility"],
        ["Generation", "answerable, claims, reason, validation errors"],
        ["Evidence", "E-ID, leaf node, quote, page/char span, support score, claim"],
        ["Result", "Toàn bộ output cùng expansion trace và metrics"],
    ])


def add_policy_knowledge(document: Document) -> None:
    document.add_heading("4. Policy thích nghi và đặc trưng", level=1)
    document.add_heading("4.1 Vector đặc trưng 14 chiều", level=2)
    add_table(document, ["Feature", "Cách hiểu"], [
        ["relevance", "Reranker score trực tiếp của ancestor candidate"],
        ["coverage", "Số member được giữ / tổng slot member"],
        ["coherence", "Jaccard trung bình giữa các member kề nhau"],
        ["density", "Member tokens / candidate tokens, chặn tại 1"],
        ["noise", "1 − density"],
        ["cost", "Token tăng thêm / context token budget"],
        ["query_* (4)", "One-hot factoid, explanatory, comparative, global"],
        ["member_count_norm", "Số member / 8, chặn tại 1"],
        ["member_score_entropy", "Entropy chuẩn hóa của phân bố member score"],
        ["section_tokens_norm", "Độ lớn section tương đối với budget"],
        ["query_length_norm", "Số từ query / 25, chặn tại 1"],
    ])
    add_formula(
        document,
        "density = min(1, Σ member_tokens / candidate_tokens)",
        "Đây là đặc trưng evidence density giữ lại từ tên gọi EDAHR ban đầu.",
    )
    document.add_heading("4.2 Calibrated prior", level=2)
    document.add_paragraph(
        "Khi không có checkpoint, AdaptiveMergePolicy dùng logistic prior với hệ số cố định: thưởng relevance, "
        "coverage, coherence, density và query toàn cục; phạt noise, cost và factoid. Prior ra xác suất để "
        "monitoring, nhưng quyết định dựa trên utility delta và margin."
    )
    document.add_heading("4.3 Learned gate", level=2)
    document.add_paragraph(
        "Checkpoint học thay utility comparison bằng xác suất P(merge) và threshold. Parent gate và section "
        "gate độc lập, có thể bật/tắt riêng để ablation. Runtime hỗ trợ TorchScript .ts và sklearn .joblib."
    )
    add_callout(
        document,
        "Ranh giới diễn giải",
        "Xác suất gate chỉ nên gọi là risk/merge probability khi calibration đã được đo bằng Brier score, ECE "
        "và reliability curve. Báo cáo V7 hiện yêu cầu hoàn thiện phần này.",
    )


def add_training_pipeline(document: Document) -> None:
    document.add_heading("5. Pipeline tạo dữ liệu và huấn luyện", level=1)
    document.add_heading("5.1 Counterfactual rollout", level=2)
    document.add_paragraph(
        "Với mỗi query và candidate group, RolloutRunner chạy các thế giới phản thực qua cùng generator và "
        "verifier để đo hệ quả thực tế của quyết định hierarchy."
    )
    add_table(document, ["Nhánh", "Can thiệp"], [
        ["KEEP", "Giữ các child/member đã truy hồi"],
        ["PARENT / EXPAND", "Thay member bằng parent chung"],
        ["SECTION", "Thay bằng section ở mức cao hơn khi tồn tại"],
    ])
    document.add_heading("5.2 Reward có ràng buộc attribution", level=2)
    add_formula(
        document,
        "R = w_a·AnswerF1 + w_r·CitationRecall + w_p·CitationPrecision + w_res·Rescue − λ_h·Harmful − λ_a·Ambiguity − λ_e·Empty − β·Tokens − γ·Latency",
        "Reward cân bằng chất lượng câu trả lời, khả năng cứu bằng chứng, drift có hại và chi phí hệ thống.",
    )
    document.add_paragraph(
        "Nhãn EXPAND chỉ dương khi precision và recall không giảm quá epsilon so với KEEP, harmful drift dưới "
        "delta, không tạo empty evidence mới và reward gain vượt tau. Với dữ liệu không có gold, code còn giữ "
        "legacy reward dựa trên coverage, citation quality và Attribution Risk."
    )
    document.add_heading("5.3 Huấn luyện và artifact", level=2)
    add_numbered(document, [
        ("Input", "JSONL rollout, mỗi row có query/source, 14 features, branch traces và hai label."),
        ("Split", "Group-disjoint theo paper source để giảm leakage."),
        ("Model", "MLP 14 → hidden 16 → ReLU → 1 logit, BCEWithLogitsLoss."),
        ("Early stopping", "Theo validation loss, giữ best state."),
        ("Export", "TorchScript checkpoint và metadata sidecar chứa SHA-256, seed, source rollout và report."),
        ("Alternative", "Tree/sklearn trainer có thể xuất .joblib cùng threshold đã chọn."),
    ])
    add_callout(
        document,
        "Tình trạng V7",
        "Fresh rollouts đã có 800 train questions/534 papers và 250 dev questions/164 papers. Tuy nhiên parent "
        "dev AUC = 0,5366 và section dev AUC = 0,5110 trong khi train gần 1,0: dấu hiệu overfit và signal yếu. "
        "Không dùng learned gate để claim superiority trước khi vượt acceptance criterion đã chốt.",
        PALE_ORANGE,
    )


def add_evaluation(document: Document) -> None:
    document.add_heading("6. Đánh giá khoa học và benchmark", level=1)
    document.add_heading("6.1 Nhóm metric", level=2)
    add_table(document, ["Nhóm", "Metric chính", "Câu hỏi đo lường"], [
        ["Retrieval", "Recall@k, Precision@k, Hit@k, MRR, nDCG", "Có tìm đúng bằng chứng và xếp đủ cao không?"],
        ["Answer", "EM, token F1, max-over-annotators", "Nội dung trả lời gần reference đến đâu?"],
        ["Grounding", "Citation P/R/F1, paragraph Evidence F1", "Citation có đúng và đủ không?"],
        ["Attribution", "AR, unsupported rate, survival, rescue/harmful drift", "Claim có bám leaf cụ thể không?"],
        ["Selective", "Risk-coverage, AURC, E-AURC", "Confidence có xếp đúng rủi ro abstention không?"],
        ["System", "Token, latency, cost, acceptance/rollback", "Chất lượng đạt được với chi phí nào?"],
        ["Statistics", "Clustered CI, paired test, effect size", "Chênh lệch có ổn định theo paper không?"],
    ])
    document.add_heading("6.2 QASPER", level=2)
    add_bullets(document, [
        "Giữ reference_answers và reference_evidence_sets riêng theo annotator.",
        "Answer EM/F1 và paragraph Evidence F1 lấy max trên annotator, theo normalization kiểu SQuAD/QASPER.",
        "Stable paragraph_id nối output child về paragraph gốc, tránh overlap chunk nhân đôi denominator.",
        "Evaluator local đã có semantics max-over-annotators, nhưng audit độc lập bằng official AllenAI script vẫn là P0.",
    ])
    document.add_heading("6.3 SciFact OOD", level=2)
    document.add_paragraph(
        "SciFact được dùng như thí nghiệm ngoài phân phối về rationale attribution. Không diễn giải answer F1 "
        "hiện tại thành stance accuracy và không gọi đây là bằng chứng tổng quát cho scientific QA nếu chưa có "
        "benchmark QA độc lập tương ứng."
    )
    document.add_heading("6.4 Baseline và ablation cần thiết", level=2)
    add_bullets(document, [
        "BM25, dense-only, hybrid/RRF, flat neural, static hierarchy, prior adaptive, learned adaptive.",
        "Parent-only, section-only, không rollback, không verifier và các ablation đúng với tên công bố.",
        "Oracle evidence/context và full-document chỉ được dùng khi được triển khai đúng nghĩa; V7 phát hiện ba label hiện có chưa đúng semantics.",
        "Main test 180 paper phải đóng cho đến khi dev acceptance, evaluator audit và statistical protocol hoàn tất.",
    ])


def add_software_architecture(document: Document) -> None:
    document.add_heading("7. Kiến trúc phần mềm và vận hành", level=1)
    document.add_heading("7.1 Bản đồ module", level=2)
    add_table(document, ["Module", "Trách nhiệm"], [
        ["ingestion.py", "Đọc PDF bằng Docling, giữ layout/page metadata"],
        ["hierarchy.py / invariants.py", "Xây và kiểm tra cây Document–Section–Parent–Child"],
        ["index.py", "BGE-M3 dense/sparse/ColBERT và fusion"],
        ["models.py", "Encoder, reranker, NLI, OpenAI/Gemini/Antigravity adapters"],
        ["policy.py / expansion.py", "Feature extraction, merge decision, rollback và iterative expansion"],
        ["context.py", "Dedup, knapsack, diversity, ContextBlock"],
        ["verification.py / attribution.py", "Claim-to-leaf NLI, safety guards và AR metrics"],
        ["pipeline.py", "Điều phối end-to-end và telemetry"],
        ["rollouts.py / training.py", "Counterfactual data và learned checkpoints"],
        ["qasper.py / scifact.py / evaluation.py", "Dataset conversion và benchmark metrics"],
        ["runtime.py / cli.py / api.py", "Dependency assembly, command line và HTTP surface"],
    ])
    document.add_heading("7.2 Dependency injection", level=2)
    document.add_paragraph(
        "Pipeline phụ thuộc các interface Retriever, Reranker, Generator và Verifier. Thiết kế này cho phép "
        "thay provider/model, chạy baseline hoặc test bằng fake adapter mà không đổi logic điều phối."
    )
    document.add_heading("7.3 Cấu hình mặc định quan trọng", level=2)
    add_table(document, ["Tham số", "Mặc định", "Vai trò"], [
        ["child_target_tokens", "220", "Kích thước leaf"],
        ["candidate_k / rerank_k", "80 / 24", "Độ rộng retrieval/rerank"],
        ["final_context_k", "8", "Số block cuối tối đa"],
        ["context_token_budget", "7000", "Ngân sách context"],
        ["merge_threshold / margin", "0,56 / 0,04", "Ngưỡng learned / prior"],
        ["rollback_ratio", "0,78", "Guard mất specificity"],
        ["nli support / contradiction", "0,25 / 0,50", "Ngưỡng verification"],
        ["claim_confidence_threshold", "0,55", "Lọc claim confidence thấp"],
        ["max_evidence_per_claim", "1", "Leaf citation tối đa mỗi claim"],
        ["seed", "42", "Khả năng tái lập"],
    ])
    document.add_heading("7.4 Luồng triển khai", level=2)
    add_numbered(document, [
        ("Build", "Load PDF → hierarchy → index → model adapters → policies."),
        ("Serve", "CLI edahr hoặc API nhận query; source filter giới hạn theo paper khi benchmark."),
        ("Observe", "Kết quả JSON giữ latency từng pha, merge/rollback, raw generation và verification trace."),
        ("Reproduce", "Config, selection manifest, checkpoint và rollout được băm SHA-256 trong artifact protocol."),
    ])


def add_knowledge_foundations(document: Document) -> None:
    document.add_heading("8. Nền tảng kiến thức xây dựng hệ thống", level=1)
    document.add_paragraph(
        "Hệ thống là giao điểm của nhiều lĩnh vực. Bảng dưới nêu kiến thức cần nắm và nơi kiến thức đó xuất "
        "hiện trong thiết kế."
    )
    add_table(document, ["Lĩnh vực", "Kiến thức", "Ứng dụng trong EDAHR"], [
        ["Scientific document processing", "Layout analysis, heading/table extraction, page geometry", "Docling ingestion và provenance trang"],
        ["Information retrieval", "Dense retrieval, learned sparse, inverted index, RRF", "Tìm leaf đa biểu diễn"],
        ["Neural ranking", "Cross-encoder relevance", "Rerank leaf và ancestor candidate"],
        ["Hierarchical representation", "Chunking, overlap, tree aggregation", "Document–Section–Parent–Child"],
        ["NLP semantics", "NLI entailment/contradiction, negation, numeric consistency", "Claim verification và safety guard"],
        ["Grounded generation", "Structured output, citation contracts, abstention", "Atomic claims có context IDs"],
        ["Decision theory", "Counterfactual action, utility, constrained optimization", "KEEP/EXPAND/SECTION policy"],
        ["Machine learning", "MLP/tree classifier, group split, calibration", "Hai learned gate độc lập"],
        ["Optimization", "0/1 knapsack, dedup, constrained diversity", "Context dưới token budget"],
        ["Evaluation science", "Multi-annotator metrics, OOD, ablation, leakage control", "QASPER/SciFact protocol"],
        ["Statistics", "Cluster bootstrap, paired tests, effect size, Holm correction", "Kết luận ở đơn vị paper"],
        ["MLOps/reproducibility", "Hashes, frozen manifests, model revisions, trace logging", "Publication artifacts và one-shot test"],
    ])
    document.add_heading("8.1 Chuỗi lập luận kỹ thuật", level=2)
    add_numbered(document, [
        ("Vấn đề", "Flat retrieval chính xác cục bộ nhưng có thể thiếu ngữ cảnh; fixed expansion đầy đủ nhưng dễ nhiễu."),
        ("Biểu diễn", "Giữ leaf làm đơn vị chân lý, ancestor làm context candidate có descendants rõ ràng."),
        ("Quyết định", "Ước lượng lợi ích mở rộng từ relevance, coverage, density, coherence, cost và query type."),
        ("Bảo vệ", "Rollback, token budget, source diversity và query ceiling giới hạn rủi ro trước generation."),
        ("Kiểm chứng", "Structured citation contract và NLI đưa mọi claim trở lại leaf."),
        ("Học", "Counterfactual rollouts biến chất lượng downstream thành nhãn cho gate."),
        ("Đánh giá", "Tách official paragraph evidence khỏi leaf attribution để tránh claim metric sai."),
    ])


def add_limitations(document: Document) -> None:
    document.add_heading("9. Giới hạn và rủi ro hiện tại", level=1)
    add_table(document, ["Mức", "Vấn đề", "Hành động bắt buộc"], [
        ["P0", "Learned gate overfit, dev gần chance", "Ổn định target, audit feature shift, calibration, đặt dev gate"],
        ["P0", "Một số baseline mang tên oracle/full-document nhưng chưa đúng nghĩa", "Triển khai đúng hoặc loại khỏi bảng"],
        ["P0", "Evaluator audit chưa chạy official AllenAI code", "Pin/vendor evaluator và yêu cầu exact fixture agreement"],
        ["P0", "Chưa có full dev acceptance benchmark", "Chạy đủ baseline/ablation sau khi protocol đóng băng"],
        ["P0", "Thống kê chưa hoàn chỉnh", "Clustered test/CI, effect size, Holm, non-inferiority"],
        ["P1", "Generator swap, SciFact OOD và manual audit chưa đủ", "Chạy matched subset và audit mù theo manifest"],
        ["P1", "Provenance/CI/license disclosure chưa hoàn chỉnh", "Bổ sung artifact README, clean CI và disclosure"],
    ])
    add_callout(
        document,
        "Quy tắc one-shot test",
        "Không chạy main test 180 paper trước khi mọi P0 đã fix và regression-tested, method/threshold/prompt/metric "
        "đã freeze, full dev đạt acceptance, và snapshot provenance đã archive. Sau test không dùng test label để sửa hệ thống.",
        PALE_ORANGE,
    )
    document.add_heading("9.1 Rủi ro kỹ thuật", level=2)
    add_bullets(document, [
        "Chi phí GPU/RAM cao do BGE-M3 multi-vector, reranker và NLI cùng tồn tại.",
        "LLM generation tạo nhiễu cho counterfactual label dù temperature bằng 0 không bảo đảm hoàn toàn deterministic.",
        "Rule-based query classification có thể phân loại sai câu hỏi ngoài từ khóa đã biết.",
        "Lexical fallback hữu ích với near-verbatim claim nhưng cần kiểm thử liên tục cho phủ định và số liệu.",
        "Docling/page mapping phụ thuộc chất lượng PDF; scan hoặc layout dị thường làm provenance kém tin cậy.",
    ])


def add_usage_and_references(document: Document) -> None:
    document.add_heading("10. Cách sử dụng và đọc output", level=1)
    document.add_heading("10.1 Chạy CLI", level=2)
    paragraph = document.add_paragraph()
    run = paragraph.add_run(
        'edahr paper-a.pdf paper-b.pdf --question "How does the method improve evidence retrieval?" --config config.json'
    )
    run.font.name = "Cascadia Mono"
    run.font.size = Pt(9)
    add_callout(
        document,
        "Bảo mật cấu hình",
        "Không commit config.local.json, API key hoặc nội dung secret. Artifact provenance chỉ lưu hash cấu hình khi cần.",
    )
    document.add_heading("10.2 Thứ tự đọc kết quả", level=2)
    add_numbered(document, [
        ("generation", "Claim cuối đã qua verification và citation E-ID."),
        ("evidence", "Quote leaf, nguồn, trang, char span và support score."),
        ("decisions", "Gate nào merge/retain/rollback và vì sao."),
        ("context", "Block C-ID đã đưa vào prompt, level và descendants."),
        ("raw_generation", "Output trước verifier cùng contract errors."),
        ("verification_trace", "NLI/contradiction/lexical score cho từng child candidate."),
        ("metrics", "Token, latency, attribution risk, rejection và acceptance rates."),
    ])

    document.add_heading("Phụ lục A. Thuật ngữ", level=1)
    add_table(document, ["Thuật ngữ", "Định nghĩa ngắn"], [
        ["Leaf / child", "Đoạn nhỏ nhất, đơn vị retrieval và citation cuối"],
        ["Parent", "Nhóm nhiều child gần nhau trong cùng section"],
        ["Expansion", "Thay các node chi tiết bằng ancestor rộng hơn"],
        ["Attribution", "Khả năng quy một claim về leaf hỗ trợ cụ thể"],
        ["Harmful drift", "Evidence mới ngoài retrieval nhưng không thuộc gold"],
        ["Rescue", "Evidence gold được expansion/verification tìm lại ngoài retrieval ban đầu"],
        ["Rollback", "Hủy merge khi ancestor mất relevance so với member tốt nhất"],
        ["RRF", "Reciprocal Rank Fusion"],
        ["NLI", "Natural Language Inference: entailment/neutral/contradiction"],
        ["AURC", "Area Under Risk-Coverage curve"],
    ])

    document.add_heading("Phụ lục B. Nguồn kỹ thuật trong repository", level=1)
    sources = [
        "README.md — mô tả pipeline và ranh giới hiện tại.",
        "AI_HANDOFF_PUBLICATION_V7.md — protocol và yêu cầu publication.",
        "analysis/v7_publication_readiness_audit.md — trạng thái và blockers ngày 01/09/2026.",
        "src/edahr/pipeline.py — orchestration trực tuyến.",
        "src/edahr/hierarchy.py, ingestion.py — cấu trúc dữ liệu và PDF provenance.",
        "src/edahr/index.py, models.py — retrieval và model adapters.",
        "src/edahr/policy.py, expansion.py, context.py — quyết định mở rộng và token budget.",
        "src/edahr/verification.py, attribution.py — claim-to-leaf verification.",
        "src/edahr/rollouts.py, training.py — phản thực, reward và learned policy.",
        "src/edahr/evaluation.py, qasper.py, scifact.py — metric và benchmark conversion.",
        "src/edahr/config.py, runtime.py, cli.py, api.py — cấu hình và vận hành.",
    ]
    add_bullets(document, sources)
    document.add_paragraph(
        "Tài liệu này mô tả code ở trạng thái workspace ngày 01/09/2026. Các con số benchmark lịch sử "
        "không được dùng như claim superiority nếu chưa vượt publication gate của V7."
    )


def build_document() -> Path:
    document = Document()
    configure_document(document)
    add_cover(document)
    add_toc(document)
    add_chapter_overview(document)
    add_online_pipeline(document)
    add_data_architecture(document)
    add_policy_knowledge(document)
    add_training_pipeline(document)
    add_evaluation(document)
    add_software_architecture(document)
    add_knowledge_foundations(document)
    add_limitations(document)
    add_usage_and_references(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)