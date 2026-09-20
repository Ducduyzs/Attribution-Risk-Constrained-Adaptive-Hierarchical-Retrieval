from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


root = Path(__file__).resolve().parents[1]
path = root / "artifacts" / "Attribution_Risk_Adaptive_Hierarchical_Retrieval_IEEE_Draft.docx"
doc = Document(path)
text = "\n".join(p.text for p in doc.paragraphs)

assert path.stat().st_size > 100_000, path.stat().st_size
assert len(doc.sections) == 2, len(doc.sections)
assert len(doc.tables) == 5, len(doc.tables)
assert "Abstract-" in text
assert "REFERENCES" in text
import re as _re
_ref_numbers = sorted(
    int(match) for match in _re.findall(r"^\[(\d+)\]", text, flags=_re.MULTILINE)
)
assert _ref_numbers == list(range(1, len(_ref_numbers) + 1)), _ref_numbers
assert 30 <= len(_ref_numbers) <= 40, len(_ref_numbers)
for _mandatory in ("ALCE", "Bohnet", "Self-RAG", "Adaptive-RAG", "Corrective Retrieval",
                   "Graph RAG"):
    assert _mandatory in text, _mandatory
assert "83 passed tests" in text
assert "[email]" not in text
assert "TODO" not in text
assert "turn" not in text.lower() or "counterfactual" in text.lower()

for section in doc.sections:
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert round(section.left_margin.inches, 2) == 0.75
    assert round(section.right_margin.inches, 2) == 0.75

cols = [section._sectPr.xpath("./w:cols")[0].get(qn("w:num")) for section in doc.sections]
assert cols == ["1", "2"], cols

for table in doc.tables:
    grid = table._tbl.tblGrid
    widths = [int(col.get(qn("w:w"))) for col in grid]
    assert sum(widths) == 4968, widths
    assert all(width > 0 for width in widths)
    for row in table.rows:
        assert len(row.cells) == len(widths)
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            assert int(tcW.get(qn("w:w"))) == width

with ZipFile(path) as archive:
    names = set(archive.namelist())
    assert "word/media/image1.png" in names
    assert "word/document.xml" in names
    assert "word/styles.xml" in names
    assert "word/numbering.xml" in names

print({
    "path": str(path),
    "bytes": path.stat().st_size,
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "sections": len(doc.sections),
    "columns": cols,
    "embedded_images": 1,
    "status": "structural-audit-passed",
})
